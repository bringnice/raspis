import io
import re
import datetime as dt

import openpyxl

GROUP = "РУП9-26БП"
SHEET_RE = re.compile(r"\d{2}\.\d{2}-\d{2}\.\d{2}")      # листы-недели: "05.10-10.10 (...)"
TEACHER = re.compile(r"[А-ЯЁ][а-яё]+(?:-[А-ЯЁ][а-яё]+)?\s+(?:[А-ЯЁ]\.\s?[А-ЯЁ]\.|[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+)")
CANCEL = re.compile(r"\d{1,2}\.\d{2}\.?\s*отмена", re.I)
DATE_PREFIX = re.compile(r"^\d{1,2}\.\d{2}\.?\s*")
NOTE = re.compile(r"ЛАБ|подгрупп", re.I)


def split_lessons(text):
    """Текст ячейки -> список пар. Пара заканчивается строкой с преподавателем."""
    lessons, cur = [], []
    for line in (l.strip() for l in text.split("\n")):
        if not line:
            continue
        if NOTE.search(line) and lessons and not cur:
            lessons[-1]["note"] = (lessons[-1]["note"] + " " + line).strip()
            continue
        cur.append(line)
        if TEACHER.search(line):
            lessons.append({"subject": " ".join(cur[:-1]), "teacher": line, "note": ""})
            cur = []
    if cur:  # текст без преподавателя
        lessons.append({"subject": " ".join(cur), "teacher": "", "note": ""})
    return lessons


def parse_cell(text):
    text = (text or "").strip()
    if not text:
        return []
    m = CANCEL.search(text)
    if m:
        lessons = split_lessons(text[m.end():])
        for i, l in enumerate(lessons):
            l["cancelled"] = i == 0          # первая — отменена, остальные — замена
            l["replacement"] = i > 0
        return lessons
    lessons = split_lessons(DATE_PREFIX.sub("", text))
    for l in lessons:
        l["cancelled"] = l["replacement"] = False
    return lessons


def room_str(v):
    if v is None:
        return ""
    return str(int(v)) if isinstance(v, float) and v.is_integer() else str(v).strip()


def parse_sheet(ws):
    col = next((c.column for c in ws[5] if str(c.value or "").strip() == GROUP), None)
    if col is None:
        return {}
    numerator = "Числитель" in ws.title        # верхняя строка ячейки = числитель
    spans = {}
    for mr in ws.merged_cells.ranges:
        for r in range(mr.min_row, mr.max_row + 1):
            spans[(r, mr.min_col)] = mr

    def val(r, c):
        mr = spans.get((r, c))
        return ws.cell(mr.min_row, mr.min_col).value if mr else ws.cell(r, c).value

    days = {}
    for r0 in range(6, ws.max_row + 1):
        d = ws.cell(r0, 1).value
        if not isinstance(d, dt.datetime):
            continue
        day = {"week": "числитель" if numerator else "знаменатель", "note": "", "lessons": []}
        for k in range(6):
            r = r0 + 2 * k
            mr = spans.get((r, col))
            if mr and mr.max_row - mr.min_row >= 5:           # «День самостоятельной подготовки»
                day["note"] = str(val(r, col)).replace("\n", " ").strip()
                break
            single = mr is not None and mr.max_row > r
            pick = r if (single or numerator) else r + 1
            text = val(pick, col)
            times = re.findall(r"\d{1,2}\.\d{2}", str(ws.cell(r, 3).value or ""))
            for l in parse_cell(text):
                l.update(n=k + 1, room=room_str(val(pick, col + 1)),
                         time=f"{times[0]}–{times[-1]}" if times else "")
                day["lessons"].append(l)
        days[d.date().isoformat()] = day
    return days


def parse_workbook(raw: bytes):
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    days = {}
    for ws in wb.worksheets:
        if SHEET_RE.search(ws.title):
            days.update(parse_sheet(ws))
    return days


if __name__ == "__main__":
    import json, sys
    with open(sys.argv[1], "rb") as f:
        print(json.dumps(parse_workbook(f.read()), ensure_ascii=False, indent=1))
